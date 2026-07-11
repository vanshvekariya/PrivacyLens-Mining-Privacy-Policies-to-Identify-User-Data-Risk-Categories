"""Generate progress_report.docx for CSC 503 (Group 6).

Formatting per assignment: Times New Roman 11pt, single column, 1-inch
margins, single-spaced. Content follows progress_report_plan.md and the
results in results/. Export to PDF (progress_report.pdf) for submission.

Style rules for this document: no em dashes anywhere; bullet lists for
enumerations (tasks, research questions, plan items); headings bold 11pt.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "progress_report.docx"
FIGS = ROOT / "results" / "figures"

doc = Document()

for section in doc.sections:
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.0
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(3)

bullet_style = doc.styles["List Bullet"]
bullet_style.font.name = "Times New Roman"
bullet_style.font.size = Pt(11)
bullet_style.paragraph_format.line_spacing = 1.0
bullet_style.paragraph_format.space_before = Pt(0)
bullet_style.paragraph_format.space_after = Pt(1)
bullet_style.paragraph_format.widow_control = False


def _runs(p, text, bold=False, italic=False, size=None):
    # minimal **bold** inline markup
    for i, chunk in enumerate(text.split("**")):
        if not chunk:
            continue
        run = p.add_run(chunk)
        run.bold = bold or (i % 2 == 1)
        run.italic = italic
        if size is not None:
            run.font.size = Pt(size)


def para(text, bold=False, italic=False, align=None, size=None,
         space_before=None, space_after=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    _runs(p, text, bold=bold, italic=italic, size=size)
    return p


def bullet(text, last=False):
    p = doc.add_paragraph(style="List Bullet")
    if last:
        p.paragraph_format.space_after = Pt(5)
    _runs(p, text)
    return p


def heading(text):
    para(text, bold=True, space_before=5, space_after=3)


def caption(text):
    para(text, italic=True, size=10, space_before=1, space_after=3)


def table(rows, widths=None, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.paragraphs[0].paragraph_format.space_after = Pt(1)
            run = cell.paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
            if header and r == 0:
                run.bold = True
            if widths:
                cell.width = Inches(widths[c])
    return t


# ---------------- Title block ----------------
para("PrivacyLens: Mining Privacy Policies to Identify User Data Risk "
     "Categories", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER,
     space_after=2)
para("CSC 503 Progress Report (Group 6)", align=WD_ALIGN_PARAGRAPH.CENTER,
     space_after=2)
para("Vansh Vikunj Vekaria  ·  Lakshita Lakshita  ·  "
     "Aishwarya Pujitha Vakulabharanam Siripurapu",
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

# ---------------- 1. The problem ----------------
heading("1. The problem")
para("Privacy policies are the main way companies disclose how they "
     "collect, use, share, and retain user data, yet they are usually too "
     "long and too complex for ordinary users to read. Survey evidence "
     "shows that only a small share of adults consistently read privacy "
     "policies before accepting them [2], and a longitudinal study of "
     "policies from 1996 to 2021 found that they have grown steadily longer "
     "and harder to read [3]. As a result, users routinely agree to terms "
     "without noticing important clauses about data sharing, advertising, "
     "retention, or deletion rights.")
para("**PrivacyLens** addresses this problem as a supervised text-"
     "classification task. Given a privacy-policy segment, the system "
     "predicts which privacy-practice categories the segment expresses and "
     "presents the prediction in a user-facing form: a simplified category, "
     "a risk level, a reading-priority recommendation, and highlighted "
     "supporting text. The goal is not legal advice but triage, helping a "
     "user decide which parts of a policy deserve attention.")
para("Our dataset is the **OPP-115 corpus** [1], which contains 115 website "
     "privacy policies segmented and annotated with ten data-practice "
     "categories by law students. We use the corpus's consolidated "
     "annotations, which yield 3,791 annotated segments, and map the ten "
     "legally precise categories to seven user-facing categories (Table 1).")
para("The problem statement is unchanged since the formal proposal, but "
     "three aspects of its formulation were refined once we started working "
     "with the data. First, the task is now explicitly **multi-label**. The "
     "proposal flagged the risk that the original annotations might not "
     "match our clause-level task exactly, and this risk materialized: "
     "49.1% of annotated segments carry more than one category (1.64 labels "
     "per segment on average), so we make one binary decision per category "
     "for each segment instead of forcing a single label. This confirms "
     "that treating the task as multi-label is necessary rather than "
     "simply a modeling choice. Second, the "
     "advertising/profiling category we originally envisioned does not "
     "exist as a top-level OPP-115 practice. Advertising appears as a "
     "purpose attribute inside the collection and sharing practices, so "
     "such clauses fall under Data Collection or Third-Party Sharing; "
     "recovering a distinct advertising category from the purpose "
     "attributes is noted as future work. Third, we added a small "
     "**generalization dataset**: because OPP-115 predates GDPR and CCPA, "
     "we will hand-label 100–200 segments from 5–10 modern (2025–2026) "
     "privacy policies and use them purely as an additional evaluation "
     "set.")

# ---------------- 2. Goals ----------------
heading("2. Goals")
para("The project is organized around three research questions:")
bullet("**RQ1:** How accurately can privacy-policy segments be classified "
       "into user-facing privacy-practice categories, and which categories "
       "are inherently harder?")
bullet("**RQ2:** How much does contextual language understanding (a "
       "fine-tuned transformer) improve over rule-based and classical "
       "bag-of-words models on this task, and why?")
bullet("**RQ3:** How well do models trained on OPP-115 generalize to the "
       "privacy policies of modern services, and which categories transfer "
       "poorly?", last=True)
para("Success is measured quantitatively. Our headline metric is "
     "**macro-averaged F1**, because the label distribution is heavily "
     "skewed (Table 1) and macro-F1 weights rare categories equally. We "
     "also report micro-F1 and per-category precision, recall, and F1, "
     "together with confusion analyses between frequently confused "
     "categories. Because a missed third-party-sharing clause defeats the "
     "purpose of the tool more than a misclassified policy-change sentence "
     "does, we track recall on high-risk categories separately. Our "
     "reporting convention: model comparison tables report the mean ± "
     "standard deviation across cross-validation folds, whereas "
     "per-category metrics and confusion analyses are computed from pooled "
     "out-of-fold predictions. Consequently, the per-category F1 values in "
     "Table 3 do not average exactly to the fold-averaged macro-F1 in "
     "Table 2.")
para("The evaluation is designed so that each possible outcome answers one "
     "of the research questions. If the "
     "transformer clearly outperforms the linear models, contextual "
     "understanding matters for hedged privacy language; if not, "
     "privacy-practice classification is largely lexical. Either result "
     "answers RQ2. The generalization experiment will likewise either "
     "support using OPP-115-trained models on current policies or quantify "
     "how much post-GDPR language has drifted.")
para("Our goals have become more specific since the proposal in three "
     "ways:")
bullet("The metrics were sharpened from generic accuracy/precision/recall/"
       "F1 to macro- and micro-F1 with per-category analysis and "
       "policy-grouped cross-validation (GroupKFold).")
bullet("Two user-facing goals extend the original interpretability goal: a "
       "literature-backed risk mapping, grounded in privacy nutrition-label "
       "research [6], ToS;DR rating criteria, and GDPR data-protection "
       "principles; and a reading-priority recommendation (1–5 stars) that "
       "combines a category's risk tier with the model's prediction "
       "confidence, so that a long policy becomes a short, prioritized "
       "reading list.")
bullet("Evidence highlighting was added because the template explanations "
       "we originally planned cannot show a user which text triggered a "
       "classification. The templates are retained as the plain-English "
       "layer on top.", last=True)

# ---------------- 3. Plan and progress-to-date ----------------
heading("3. Plan and progress-to-date")
para("**Sources consulted.** As promised in the proposal, we consulted the "
     "literature before running experiments: the OPP-115 corpus paper [1], "
     "early work on automatic privacy-policy classification [7], Polisis "
     "[4], PrivBERT [5], and privacy nutrition labels [6]. These shaped "
     "concrete decisions: the multi-label framing, the choice of a "
     "transformer as the advanced model, the grounding of the risk mapping, "
     "and positioning our contribution as the user-facing layer (risk, "
     "reading priority, evidence highlighting) rather than taxonomy labels "
     "alone.")
para("**Accomplished so far** (the first three stages of the proposal "
     "timeline):")
bullet("Acquired and parsed OPP-115, completed exploratory analysis, and "
       "fixed the 10-to-7 category mapping with counts (Table 1).")
bullet("Created a policy-disjoint 80/20 train/test split (92/23 policies; "
       "3,053/738 segments; fixed seed), keeping all segments of a policy "
       "in one split to prevent document-level leakage. The test set stays "
       "untouched until the final report.")
bullet("Built a shared evaluation harness using 5-fold policy-grouped "
       "cross-validation on the training split.")
bullet("Trained and evaluated five approaches: a majority-class baseline, a "
       "keyword baseline, and TF-IDF (word unigrams and bigrams) with "
       "Multinomial Naive Bayes, Logistic Regression, and Linear SVM (the "
       "latter two class-weighted).")
bullet("Completed a first pass over the misclassified segments. Results "
       "appear in Section 5.", last=True)
para("**Remaining experiments:**")
bullet("DistilBERT (or RoBERTa-base if compute allows) fine-tuning with a "
       "multi-label classification head.")
bullet("Structured error analysis organized into recurring failure types.")
bullet("Evidence highlighting: top-weighted TF-IDF n-grams for the linear "
       "models and a simple word-occlusion score for the transformer.")
bullet("Risk mapping and the reading-priority rule.")
bullet("Labeling and evaluation of the modern-policy generalization set.")
bullet("The prototype web interface.")
bullet("Optional zero-/few-shot LLM comparison, attempted only after the "
       "core analysis is complete. Our fixed policy for spare time is depth "
       "of analysis over breadth of models.", last=True)
para("**Plan changes since the proposal:**")
bullet("Decision Trees and Random Forests, which the proposal listed, were "
       "removed: sparse, high-dimensional TF-IDF features suit linear "
       "models better, and dropping them lets us focus on the more "
       "informative comparison between linear models and a transformer.")
bullet("The proposal's generic “neural networks” item was made concrete as "
       "DistilBERT/RoBERTa fine-tuning.")
bullet("The planned fixed validation split was replaced with policy-grouped "
       "cross-validation on the training split, which uses the 115 policies "
       "more efficiently while keeping the test set clean.")
bullet("**Team change:** Mansoor Ali has left the team, leaving three "
       "members. His responsibilities (tree-based models, the evaluation "
       "harness, and error analysis) were redistributed as described in "
       "Section 4. Dropping the tree models absorbed most of the lost "
       "modeling capacity.", last=True)
para("**Updated timeline:**")
bullet("July 10–14: finalize the classical analysis and begin transformer "
       "fine-tuning.")
bullet("July 14–19: transformer experiments and structured error analysis.")
bullet("July 19–26: evidence highlighting, risk/priority layer, "
       "modern-policy labeling, and the generalization evaluation.")
bullet("July 26–31: prototype, final report, and presentation; the optional "
       "LLM baseline only if time remains.", last=True)

# ---------------- 4. Task breakdown ----------------
heading("4. Task breakdown")
para("Tasks are assigned by name; cross-cutting tasks have a named owner "
     "plus named contributors. Tasks marked (done) are complete.")
para("**Vansh Vikunj Vekaria:**", space_after=2)
bullet("Dataset understanding and label mapping (done).")
bullet("Text preprocessing and the policy-disjoint split (done).")
bullet("Linear SVM experiments (done).")
bullet("Prototype development.")
bullet("Coordination of the generalization study: policy selection and his "
       "share of the labeling.", last=True)
para("**Lakshita Lakshita:**", space_after=2)
bullet("Majority and keyword baselines (done).")
bullet("Naive Bayes experiments (done).")
bullet("Shared evaluation harness (metrics, cross-validation, confusion "
       "analyses), taken over from Mansoor (done).")
bullet("Risk-tier framework, explanation templates, and the "
       "reading-priority rule.")
bullet("Her share of the generalization labeling.", last=True)
para("**Aishwarya Pujitha Vakulabharanam Siripurapu:**", space_after=2)
bullet("TF-IDF feature extraction (done).")
bullet("Logistic Regression experiments (done).")
bullet("DistilBERT/RoBERTa fine-tuning (the concrete form of her original "
       "neural-network task).")
bullet("Coordination of the error analysis, taken over from Mansoor: each "
       "member reviews the errors of the models they own, and Aishwarya "
       "consolidates the failure-type taxonomy.")
bullet("Visualization and report assembly; her share of the generalization "
       "labeling.", last=True)
para("**Changes since the proposal:** Mansoor Ali owned the Decision Tree "
     "and Random Forest experiments, model evaluation, and error analysis. "
     "The tree models were removed for the reason given in Section 3; the "
     "evaluation harness moved to Lakshita and error-analysis coordination "
     "to Aishwarya, as reflected above.")

# ---------------- 5. Initial results ----------------
heading("5. Initial results")
para("**Dataset.** Table 1 summarizes the corpus after parsing and mapping: "
     "3,791 annotated segments across 115 policies (33 segments per policy "
     "on average, mean length 70 words). Two findings directly affected the "
     "design. First, 49.1% of segments carry more than one category, which "
     "confirms the multi-label, one-decision-per-category framing. Second, "
     "the class imbalance is severe: Data Retention has 156 segments while "
     "Other/Unclear has 1,985. This motivates class weighting and macro-F1 "
     "as the headline metric.")

table([
    ["Original OPP-115 category (segments)", "User-facing category",
     "Segments"],
    ["First Party Collection/Use (1,521)", "Data Collection", "1,521"],
    ["Third Party Sharing/Collection (1,186)", "Third-Party Sharing",
     "1,186"],
    ["User Choice/Control (632); User Access, Edit and Deletion (231)",
     "User Control & Deletion", "793"],
    ["Data Security (375)", "Data Security", "375"],
    ["Policy Change (192)", "Policy Change", "192"],
    ["Data Retention (156)", "Data Retention", "156"],
    ["Other (1,762); International and Specific Audiences (353); "
     "Do Not Track (32)", "Other/Unclear", "1,985"],
], widths=[3.4, 1.9, 0.9])
caption("Table 1: Category mapping and label distribution (all 3,791 "
        "segments). Merged categories count fewer segments than the sum of "
        "their sources because a segment carrying two source labels is "
        "counted once.")

para("**Model comparison.** Table 2 reports 5-fold policy-grouped "
     "cross-validation on the training split; the held-out test set has "
     "not yet been used and is reserved exclusively for the final "
     "evaluation.")

table([
    ["Model", "Macro-F1 (mean ± std)", "Micro-F1 (mean ± std)"],
    ["Majority class", "0.099 ± 0.004", "0.401 ± 0.030"],
    ["Multinomial Naive Bayes", "0.331 ± 0.008", "0.624 ± 0.011"],
    ["Keyword baseline", "0.530 ± 0.020", "0.556 ± 0.009"],
    ["Logistic Regression (TF-IDF)", "0.713 ± 0.024", "0.764 ± 0.008"],
    ["Linear SVM (TF-IDF)", "0.713 ± 0.024", "0.768 ± 0.007"],
], widths=[2.6, 1.8, 1.8])
caption("Table 2: Model comparison, 5-fold policy-grouped cross-validation "
        "on the training split.")

para("Three observations emerge from Table 2:")
bullet("**Logistic Regression and Linear SVM achieved essentially identical "
       "performance** (macro-F1 0.713 for both). Both are strong linear "
       "baselines. We use the SVM for the detailed analysis below only "
       "because its macro-F1 was numerically higher by roughly 0.00002, a "
       "difference with no practical meaning; we do not claim it "
       "outperformed Logistic Regression.")
bullet("Both linear models beat the keyword baseline by about 0.18 "
       "macro-F1, which is evidence that they capture patterns beyond the "
       "obvious vocabulary. This is a first partial answer to RQ2.")
bullet("**Naive Bayes underperformed the keyword baseline on macro-F1** "
       "despite a higher micro-F1. This is expected rather than an "
       "implementation error: Multinomial Naive Bayes does not support "
       "class weighting in the implementation used, so under severe "
       "imbalance its estimates favor frequent categories and recall on "
       "rare categories (Data Retention, Policy Change) collapses, which "
       "is exactly what macro-F1 penalizes. The contrast with the "
       "class-weighted linear models shows the value of imbalance "
       "handling.", last=True)

para("**Per-category results.** Table 3 and Figure 1 break down the "
     "selected model's behavior (pooled out-of-fold predictions).")

table([
    ["Category", "Precision", "Recall", "F1", "Support"],
    ["Data Collection", "0.819", "0.807", "0.813", "1,244"],
    ["Third-Party Sharing", "0.800", "0.770", "0.785", "959"],
    ["Other/Unclear", "0.817", "0.752", "0.784", "1,619"],
    ["Policy Change", "0.819", "0.739", "0.777", "153"],
    ["User Control & Deletion", "0.700", "0.703", "0.702", "637"],
    ["Data Security", "0.782", "0.601", "0.680", "293"],
    ["Data Retention", "0.658", "0.375", "0.478", "128"],
], widths=[2.2, 1.0, 1.0, 1.0, 1.0])
caption("Table 3: Linear SVM per-category precision/recall/F1 on pooled "
        "out-of-fold predictions. Support counts cover the training split "
        "only (3,053 segments), so they are lower than the full-corpus "
        "counts in Table 1.")

fig_p = doc.add_paragraph()
fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
fig_p.add_run().add_picture(str(FIGS / "fig_confusion.png"),
                            width=Inches(3.8))
caption("Figure 1: For segments with exactly one gold label, the rate at "
        "which each category is predicted (Linear SVM, out-of-fold). "
        "Off-diagonal mass shows which categories are confused; note the "
        "small sample for Data Retention (n=6 single-label segments).")

para("The high-risk categories users most need (Data Collection, "
     "Third-Party Sharing) are also the best-detected (F1 0.81 and 0.79), "
     "with recall 0.81 and 0.77 respectively. The clear weak spot is **Data "
     "Retention** (F1 0.478, recall 0.375), the rarest category, followed "
     "by Data Security's recall of 0.601. Figure 1 adds two patterns: gold "
     "User Control & Deletion segments are often also predicted "
     "Other/Unclear (0.38), and gold Third-Party Sharing segments are often "
     "co-predicted Data Collection (0.22), reflecting advertising-related "
     "language that genuinely straddles both practices.")
para("**Error examples.** Two recurring failure types from the first-pass "
     "review illustrate why privacy language is hard. (i) Cross-references: "
     "“reddit only uses and discloses your information … that we describe "
     "in the How We Use or Disclose Collected Information section” is gold "
     "Data Collection + Third-Party Sharing but was predicted "
     "Other/Unclear; the practice is described elsewhere in the document, "
     "beyond a segment-level model's view. (ii) Semantically plausible "
     "over-prediction: "
     "“we may share non-personally identifiable information … with "
     "interested third parties” is gold Third-Party Sharing only, but the "
     "model added Data Collection. Some errors of this second type are "
     "artifacts of simplifying ten legal categories into seven user-facing "
     "ones: a prediction can be semantically defensible yet count as wrong "
     "against the mapped gold label. We acknowledge this as a limitation of "
     "the mapping without using it to excuse genuine misclassifications.")
para("For completeness: under the strict exact-match criterion, 1,606 of "
     "3,053 training segments have at least one incorrect binary decision "
     "among the seven categories. This subset accuracy (47%) is far "
     "stricter than micro-F1 (0.768): one wrong label among seven marks "
     "the whole segment wrong. Such values are expected in multi-label "
     "classification and should not be read as the model being wrong half "
     "the time.")
para("**How the results affected the plan:**")
bullet("The multi-label prevalence validated the one-vs-rest framing "
       "before the transformer work begins.")
bullet("The rare-category weakness sharpens RQ2: the transformer's main "
       "test is whether context helps where lexical features fail (Data "
       "Retention, Data Security recall), so the error analysis will focus "
       "there.")
bullet("Logistic Regression and Linear SVM performed near-identically, so "
       "either is a reasonable classical choice. Logistic Regression "
       "offers directly interpretable weights for the planned evidence "
       "highlighting; Linear SVM remains the model analyzed in this "
       "report.")
bullet("The keyword baseline's respectable macro-F1 (0.530) confirms a "
       "strong lexical signal, supporting our expectation that "
       "n-gram-based evidence highlighting will be informative to users.",
       last=True)

# ---------------- References ----------------
heading("References")
refs = [
    "[1] S. Wilson, F. Schaub, A. A. Dara, F. Liu, S. Cherivirala, P. G. "
    "Leon, M. S. Andersen, S. Zimmeck, K. M. Sathyendra, N. C. Russell, "
    "T. B. Norton, E. Hovy, J. Reidenberg, and N. Sadeh. The Creation and "
    "Analysis of a Website Privacy Policy Corpus. In Proceedings of ACL, "
    "2016.",
    "[2] Pew Research Center. Americans' Attitudes and Experiences with "
    "Privacy Policies and Laws. 2019.",
    "[3] I. Wagner. Privacy Policies Across the Ages: Content and "
    "Readability of Privacy Policies 1996–2021. 2022.",
    "[4] H. Harkous, K. Fawaz, R. Lebret, F. Schaub, K. G. Shin, and K. "
    "Aberer. Polisis: Automated Analysis and Presentation of Privacy "
    "Policies Using Deep Learning. In Proceedings of USENIX Security, 2018.",
    "[5] M. Srinath, S. Wilson, and C. L. Giles. Privacy at Scale: "
    "Introducing the PrivaSeer Corpus of Web Privacy Policies. In "
    "Proceedings of ACL-IJCNLP, 2021.",
    "[6] P. G. Kelley, J. Bresee, L. F. Cranor, and R. W. Reeder. A "
    "“Nutrition Label” for Privacy. In Proceedings of SOUPS, 2009.",
    "[7] F. Liu, S. Wilson, P. Story, S. Zimmeck, and N. Sadeh. Towards "
    "Automatic Classification of Privacy Policy Text. Technical report, "
    "School of Computer Science, Carnegie Mellon University, 2018.",
]
for r in refs:
    para(r, space_after=3)

doc.save(OUT)

# guard: the document must contain no em dashes
from docx import Document as _D
check = _D(OUT)
texts = [p.text for p in check.paragraphs]
for t in check.tables:
    for row in t.rows:
        for cell in row.cells:
            texts.append(cell.text)
bad = [t for t in texts if "—" in t]
if bad:
    raise SystemExit(f"EM DASH FOUND in {len(bad)} paragraph(s): {bad[:2]}")
print("wrote", OUT, "(no em dashes)")
